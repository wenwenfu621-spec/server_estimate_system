"""
report_export.py - 報價單 Excel / Word 輸出模組
Version: v3.0.0_20261023
Description: 依 template.xlsm（工作表 Main）產出板金報價 Excel（.xlsx，不含巨集）
             與 Word 圖文報告。

Excel 欄位（第 3–4 列為表頭，資料自第 5 列起）：
    A 序項  B 示意圖  C 品名  D 材質  E 料厚  F 尺寸
    G 噴漆  H 陽極  I 電鍍  J 金額  K 利潤（K4 = 利潤 %）  L 備註
    M 長  N 寬  O 密度  P 重量(kg)  Q 單價(元/kg)   ← 輔助欄，輸出時隱藏

公式（插入列後由 Excel 自動調整參照）：
    A = IF(C="","",COUNTA($C$5:C))                    序號依品名自動重排
    P = 長 × 寬 × 料厚 × 密度 ÷ 1,000,000              重量 kg
    J = 重量 × 單價 + 噴漆 + 陽極 + 電鍍               金額（不含利潤）
    K = J × (1 + $K$4)                               含利潤售價
"""

import os
import math
import unicodedata
from copy import copy
from datetime import datetime
from typing import Dict, Any, List, Optional

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

from cad_parser import is_valid_image

_MODULE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_CANDIDATES = ["template.xlsm", "template.xlsx"]

SHEET_NAME = "Main"
FIRST_DATA_ROW = 5
TEMPLATE_DATA_ROWS = 20          # 範本內已畫好框線的資料列數（第 5–24 列）
PROFIT_CELL = "K4"
TREATMENTS = [("paint", "噴漆"), ("anodize", "陽極"), ("plating", "電鍍")]

# 欄位編號
COL_SEQ, COL_IMG, COL_NAME, COL_MAT, COL_THK, COL_DIM = 1, 2, 3, 4, 5, 6
COL_PAINT, COL_ANODIZE, COL_PLATING, COL_AMOUNT, COL_PROFIT, COL_REMARK = 7, 8, 9, 10, 11, 12
COL_LEN, COL_WID, COL_DENS, COL_WEIGHT, COL_PRICE = 13, 14, 15, 16, 17
LAST_VISIBLE_COL = COL_REMARK
HELPER_HEADERS = {COL_LEN: "長", COL_WID: "寬", COL_DENS: "密度", COL_WEIGHT: "重量(kg)", COL_PRICE: "單價(元/kg)"}

IMAGE_ROW_HEIGHT_PT = 67.5       # ≈ 90 px
IMAGE_PADDING_PX = 4


def _num(value) -> Optional[float]:
    try:
        if value is None or (isinstance(value, float) and math.isnan(value)) or value == "":
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def compute_row(item: Dict[str, Any], profit_pct: float) -> Dict[str, Optional[float]]:
    """依 Excel 相同規則算出重量、金額、含利潤售價（供網頁預覽與 Word 使用）"""
    length, width = _num(item.get("length")), _num(item.get("width"))
    thickness, density = _num(item.get("thickness")), _num(item.get("density"))
    price = _num(item.get("unit_price")) or 0.0
    weight = None
    if None not in (length, width, thickness, density):
        weight = length * width * thickness * density / 1_000_000
    fees = sum(_num(item.get(key)) or 0.0 for key, _ in TREATMENTS)
    amount = (weight or 0.0) * price + fees
    return {"weight": weight, "amount": amount, "with_profit": amount * (1 + profit_pct / 100.0)}


def _find_template() -> Optional[str]:
    for name in TEMPLATE_CANDIDATES:
        path = os.path.join(_MODULE_DIR, name)
        if os.path.exists(path):
            return path
    return None


def _display_width(text: str) -> float:
    """估算字串在 Excel 中的欄寬（全形字算 2 個字元）"""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in str(text))


def _col_width_px(ws, col: int) -> int:
    width = ws.column_dimensions[get_column_letter(col)].width or ws.sheet_format.defaultColWidth or 8.43
    return int(width * 7 + 5)


def _copy_style(src, dst):
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.fill = copy(src.fill)
    dst.protection = copy(src.protection)


def _add_cell_image(ws, image_path: str, row: int, col: int):
    """把圖片等比縮放後置中放進儲存格，並設定「隨儲存格移動和調整大小」，
    篩選或隱藏該列時圖片會跟著隱藏。"""
    cell_w = _col_width_px(ws, col)
    cell_h = int(IMAGE_ROW_HEIGHT_PT * 96 / 72)
    with Image.open(image_path) as im:
        iw, ih = im.size
    scale = min((cell_w - 2 * IMAGE_PADDING_PX) / iw, (cell_h - 2 * IMAGE_PADDING_PX) / ih)
    w, h = max(1, int(iw * scale)), max(1, int(ih * scale))
    x0, y0 = (cell_w - w) // 2, (cell_h - h) // 2

    img = XLImage(image_path)
    img.width, img.height = w, h
    anchor = TwoCellAnchor(editAs="twoCell")
    anchor._from = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(x0), row=row - 1, rowOff=pixels_to_EMU(y0))
    anchor.to = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(x0 + w), row=row - 1, rowOff=pixels_to_EMU(y0 + h))
    img.anchor = anchor
    ws.add_image(img)


def generate_excel_report(items: List[Dict[str, Any]], output_path: str, profit_pct: float = 0.0):
    """產出報價 Excel。

    items 每筆欄位：file_name, material, thickness, dimensions_str, length, width,
    density, unit_price, paint, anodize, plating, remark, iso_image_path。
    profit_pct：利潤百分比（20 代表 20%），寫入 K4。
    """
    template = _find_template()
    if template:
        # 不保留巨集：輸出一律為 .xlsx
        wb = openpyxl.load_workbook(template, keep_vba=False)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME

    # 清掉範本殘留：M 欄以後的舊公式（舊版拆解尺寸字串用）
    for row in ws.iter_rows(min_col=COL_LEN):
        for cell in row:
            if cell.value is not None:
                cell.value = None

    n = len(items)
    slots = max(n, TEMPLATE_DATA_ROWS)
    last_row = FIRST_DATA_ROW + slots - 1
    total_row = last_row + 1
    style_row = FIRST_DATA_ROW
    base_height = ws.row_dimensions[style_row].height

    # 所有資料列（含超過範本列數時補的列）與合計列，都統一套第 5 列的格式
    # （範本各列的對齊與字型原本不一致）
    for r in range(FIRST_DATA_ROW + 1, total_row + 1):
        for c in range(1, LAST_VISIBLE_COL + 1):
            _copy_style(ws.cell(style_row, c), ws.cell(r, c))
        ws.row_dimensions[r].height = base_height

    # 利潤 %
    profit_cell = ws[PROFIT_CELL]
    profit_cell.value = (profit_pct or 0.0) / 100.0
    profit_cell.number_format = "0%"

    # 輔助欄表頭（之後隱藏）
    for col, label in HELPER_HEADERS.items():
        ws.cell(4, col, label)

    widths = {c: _display_width(ws.cell(3, c).value or "") for c in range(1, LAST_VISIBLE_COL + 1)}

    for i in range(slots):
        r = FIRST_DATA_ROW + i
        item = items[i] if i < n else None

        for c in range(COL_IMG, COL_PRICE + 1):
            if c not in (COL_AMOUNT, COL_PROFIT, COL_WEIGHT):
                ws.cell(r, c).value = None

        ws.cell(r, COL_SEQ).value = f'=IF(C{r}="","",COUNTA($C$5:C{r}))'
        ws.cell(r, COL_WEIGHT).value = f'=IF(OR(M{r}="",N{r}="",E{r}="",O{r}=""),"",M{r}*N{r}*E{r}*O{r}/1000000)'
        ws.cell(r, COL_AMOUNT).value = f'=IF(C{r}="","",N(P{r})*N(Q{r})+SUM(G{r}:I{r}))'
        ws.cell(r, COL_PROFIT).value = f'=IF(J{r}="","",J{r}*(1+$K$4))'
        ws.cell(r, COL_WEIGHT).number_format = "0.000"
        ws.cell(r, COL_THK).number_format = "0.0#"
        ws.cell(r, COL_PRICE).number_format = "#,##0.00"
        for c in (COL_PAINT, COL_ANODIZE, COL_PLATING):
            ws.cell(r, c).number_format = "#,##0"
        for c in (COL_AMOUNT, COL_PROFIT):
            ws.cell(r, c).number_format = "#,##0"

        if not item:
            continue

        values = {
            COL_NAME: item.get("file_name", ""),
            COL_MAT: item.get("material") or None,
            COL_THK: _num(item.get("thickness")),
            COL_DIM: item.get("dimensions_str", ""),
            COL_PAINT: _num(item.get("paint")),
            COL_ANODIZE: _num(item.get("anodize")),
            COL_PLATING: _num(item.get("plating")),
            COL_REMARK: item.get("remark") or None,
            COL_LEN: _num(item.get("length")),
            COL_WID: _num(item.get("width")),
            COL_DENS: _num(item.get("density")) if item.get("material") else None,
            COL_PRICE: _num(item.get("unit_price")) if item.get("material") else None,
        }
        for col, val in values.items():
            ws.cell(r, col).value = val
            if val is not None and col <= LAST_VISIBLE_COL:
                widths[col] = max(widths.get(col, 0), _display_width(val))

        iso = item.get("iso_image_path")
        if iso and is_valid_image(iso):
            ws.row_dimensions[r].height = IMAGE_ROW_HEIGHT_PT
            _add_cell_image(ws, iso, r, COL_IMG)

    # 合計列
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=COL_PLATING)
    total_label = ws.cell(total_row, 1)
    total_label.value = "合  計"
    total_label.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(total_row, COL_AMOUNT).value = f"=SUM(J{FIRST_DATA_ROW}:J{last_row})"
    ws.cell(total_row, COL_PROFIT).value = f"=SUM(K{FIRST_DATA_ROW}:K{last_row})"
    for c in (COL_AMOUNT, COL_PROFIT):
        ws.cell(total_row, c).number_format = "#,##0"

    # 欄寬：依內容放寬（只加寬不縮窄），字級 12 比預設 11 大，乘上 1.15 預留
    for col, chars in widths.items():
        if col == COL_IMG:
            continue
        letter = get_column_letter(col)
        current = ws.column_dimensions[letter].width or 8.43
        ws.column_dimensions[letter].width = max(current, math.ceil(chars * 1.15) + 2)
    for col in (COL_AMOUNT, COL_PROFIT):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, 14)

    # 隱藏輔助欄 M–Q
    for col in HELPER_HEADERS:
        ws.column_dimensions[get_column_letter(col)].hidden = True

    ws.auto_filter.ref = f"A4:{get_column_letter(LAST_VISIBLE_COL)}{last_row}"
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    ws.print_area = f"A1:{get_column_letter(LAST_VISIBLE_COL)}{total_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(output_path)


def _set_run_font(run, size: float = 11, bold: bool = False):
    run.font.name = "標楷體"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")


def _fmt(value: Optional[float], digits: int = 0) -> str:
    if value is None:
        return "-"
    return f"{value:,.{digits}f}"


def generate_word_report(items: List[Dict[str, Any]], output_path: str, profit_pct: float = 0.0):
    """建立 Word 圖文報告：報價總表 + 每個零件的明細與四視圖（keep_with_next 防止跨頁拆開）"""
    if not HAS_DOCX:
        raise ModuleNotFoundError("系統缺少 python-docx 套件，無法產生 Word 報表。")

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title_p.add_run("板金件報價明細"), size=20, bold=True)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(sub_p.add_run(
        f"報價日期：{datetime.now().strftime('%Y年%m月%d日')}　　利潤：{profit_pct:g}%"), size=10)

    rows = [(item, compute_row(item, profit_pct)) for item in items]

    headers = ["序項", "品名", "材質", "料厚", "尺寸", "重量(kg)", "表面處理", "金額", "含利潤"]
    col_widths = [0.45, 2.0, 0.6, 0.5, 0.95, 0.65, 0.95, 0.6, 0.6]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        _set_run_font(cell.paragraphs[0].add_run(text), size=9, bold=True)

    total_amount = total_profit = 0.0
    for idx, (item, calc) in enumerate(rows, start=1):
        treat = "、".join(f"{label}{_fmt(_num(item.get(key)))}" for key, label in TREATMENTS if _num(item.get(key)))
        values = [
            str(idx), item.get("file_name", ""), item.get("material") or "-",
            _fmt(_num(item.get("thickness")), 2) if _num(item.get("thickness")) is not None else "-",
            item.get("dimensions_str", ""), _fmt(calc["weight"], 3), treat or "-",
            _fmt(calc["amount"]), _fmt(calc["with_profit"]),
        ]
        total_amount += calc["amount"]
        total_profit += calc["with_profit"]
        for cell, text in zip(table.add_row().cells, values):
            cell.text = ""
            _set_run_font(cell.paragraphs[0].add_run(text), size=9)

    total_cells = table.add_row().cells
    merged = total_cells[0].merge(total_cells[6])
    merged.text = ""
    _set_run_font(merged.paragraphs[0].add_run("合  計"), size=9, bold=True)
    for cell, value in ((total_cells[7], total_amount), (total_cells[8], total_profit)):
        cell.text = ""
        _set_run_font(cell.paragraphs[0].add_run(_fmt(value)), size=9, bold=True)
    for column, width in zip(table.columns, col_widths):
        column.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, col_widths):
            cell.width = Inches(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    for idx, (item, calc) in enumerate(rows, start=1):
        p_item = doc.add_paragraph()
        p_item.paragraph_format.keep_with_next = True
        _set_run_font(p_item.add_run(f"【項目 {idx}】 {item.get('file_name', '')}"), size=12, bold=True)

        unit = item.get("unit", "mm")
        thickness = _num(item.get("thickness"))
        lines = [
            f"• 材質：{item.get('material') or '未選擇'}"
            + (f"（密度 {_fmt(_num(item.get('density')), 2)} g/cm³，單價 {_fmt(_num(item.get('unit_price')), 2)} 元/kg）"
               if item.get("material") else ""),
            f"• 料厚：{_fmt(thickness, 2) + ' ' + unit if thickness is not None else '未判定'}",
            f"• 素材尺寸 (長*寬*高)：{item.get('dimensions_str', '')} {unit}（{item.get('used_mode', 'OBB')} 模式）",
            f"• 重量（長×寬×料厚×密度）：{_fmt(calc['weight'], 3)} kg",
            f"• 金額：{_fmt(calc['amount'])} 元　含利潤：{_fmt(calc['with_profit'])} 元",
        ]
        if item.get("remark"):
            lines.append(f"• 備註：{item['remark']}")
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.left_indent = Inches(0.2)
        p_desc.paragraph_format.keep_with_next = True
        _set_run_font(p_desc.add_run("\n".join(lines)), size=11)

        img_path = item.get("image_path")
        if is_valid_image(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(img_path, width=Inches(6.2))
        else:
            err = item.get("image_error")
            p_none = doc.add_paragraph(f"   [ 無法產生視圖預覽{f'（{err}）' if err else ''} ]")
            p_none.runs[0].font.color.rgb = RGBColor(128, 128, 128)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    doc.save(output_path)
